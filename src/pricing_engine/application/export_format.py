import logging
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border
from openpyxl.utils import get_column_letter

from configs.settings import STAGING_DIR, OUTPUT_DIR
from configs.constants import (
    SPECIAL_OFFER_LIQUIDITY,
    EXCLUDE_FLAGS_FROM_SPECIAL_OFFER,
    SHEET_ALL,
    SHEET_HYUNDAI_KIA,
    SHEET_SPECIAL,
    HYUNDAI_KIA_PRICE_GROUP,
)

logger = logging.getLogger(__name__)

TODAY_STR = datetime.today().strftime("%Y_%m_%d")

TECH_FILE = STAGING_DIR / "corrected_price_6.xlsx"
FINAL_FILE = OUTPUT_DIR / f"Лига-М_запчасти_{TODAY_STR}.xlsx"

DISPLAY_COLUMNS = [
    "price_group",
    "article",
    "name",
    "stock",
    "unit",
    "price",
    "flag",
]

RUS_HEADERS = [
    "Ценовая группа",
    "Артикул",
    "Номенклатура",
    "Остаток",
    "Ед.",
    "Цена (с НДС)",
    "",
]


def auto_width(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is not None and cell.value != "":
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = max_length + 2


def format_sheet(ws, df, title):
    ws.title = title

    # Header block
    ws.merge_cells("A1:G1")
    ws.merge_cells("A3:G3")
    ws.merge_cells("A5:G5")

    ws["A1"] = "Прайс-лист"
    ws["A1"].font = Font(name="Arial", size=36, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = "ООО Лига-М"
    ws["A3"].font = Font(name="Arial", size=14)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A5"] = "parts@liga-m.pro"
    ws["A5"].font = Font(name="Arial", size=8)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    # Header row (row 7)
    for col, header in enumerate(RUS_HEADERS, start=1):
        cell = ws.cell(row=7, column=col, value=header)
        cell.font = Font(name="Arial", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border()

    if df is None or df.empty:
        auto_width(ws)
        return

    # Data rows start from row 8 (no duplicated English headers)
    for r_idx, row in enumerate(df.itertuples(index=False), start=8):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Number formats + New highlight
    for row in ws.iter_rows(min_row=8):
        # D (stock)
        if row[3].value is not None:
            row[3].number_format = "#,##0"

        # F (price)
        if row[5].value is not None:
            row[5].number_format = "#,##0.00"

        # G (flag) — red, not bold
        if len(row) >= 7 and row[6].value == "New!":
            row[6].font = Font(color="FF0000")

    auto_width(ws)


def main():
    logger.info("START → export_format")

    if not TECH_FILE.exists():
        raise FileNotFoundError(f"{TECH_FILE} not found")

    df = pd.read_excel(TECH_FILE)

    missing = [c for c in DISPLAY_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    # All parts sheet
    df_all = df[DISPLAY_COLUMNS].copy()

    # Hyundai/Kia sheet
    if "price_group" in df.columns:
        df_hk = df[df["price_group"] == HYUNDAI_KIA_PRICE_GROUP][DISPLAY_COLUMNS].copy()
    else:
        df_hk = pd.DataFrame(columns=DISPLAY_COLUMNS)

    # Special offer sheet (dead+low) excluding flagged items (e.g., New!)
    if "liquidity_category" in df.columns:
        liquidity = df["liquidity_category"].astype(str).str.strip().str.lower()

        # normalize flags safely
        flags = df["flag"].fillna("").astype(str).str.strip()

        exclude_flags = set([str(x).strip() for x in EXCLUDE_FLAGS_FROM_SPECIAL_OFFER])
        special_liq = set([str(x).strip().lower() for x in SPECIAL_OFFER_LIQUIDITY])

        mask = liquidity.isin(special_liq) & (~flags.isin(exclude_flags))

        df_special = df[mask][DISPLAY_COLUMNS].copy()
    else:
        df_special = pd.DataFrame(columns=DISPLAY_COLUMNS)

    wb = Workbook()

    ws_all = wb.active
    format_sheet(ws_all, df_all, SHEET_ALL)

    ws_hk = wb.create_sheet()
    format_sheet(ws_hk, df_hk, SHEET_HYUNDAI_KIA)

    ws_special = wb.create_sheet()
    format_sheet(ws_special, df_special, SHEET_SPECIAL)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(FINAL_FILE)

    logger.info(f"Saved formatted file: {FINAL_FILE.name}")
    logger.info("DONE → export_format")


if __name__ == "__main__":
    main()